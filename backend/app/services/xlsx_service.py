import io
from dataclasses import dataclass
from datetime import date, datetime
from typing import Literal

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ColumnKind = Literal["text", "long_text", "integer", "number", "money", "datetime"]
CellValue = str | int | float | datetime | date | None

MONEY_FORMAT = '"CHF" #,##0.00'
NUMBER_FORMAT = "#,##0.##"
INTEGER_FORMAT = "0"
DATETIME_FORMAT = "DD.MM.YYYY HH:MM"
MIN_WIDTH = 8
MAX_WIDTH = 60
LONG_TEXT_WIDTH = 80
FIXED_WIDTHS: dict[ColumnKind, int] = {"money": 16, "datetime": 17}
HEADER_FONT = Font(bold=True)
TOP_ALIGNMENT = Alignment(vertical="top")
WRAPPED_ALIGNMENT = Alignment(vertical="top", wrap_text=True)


@dataclass(frozen=True)
class XlsxColumn:
    header: str
    kind: ColumnKind


@dataclass(frozen=True)
class XlsxSheet:
    title: str
    columns: list[XlsxColumn]
    rows: list[list[CellValue]]


def _clean_text(value: str) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("", value)


def _write_value(cell: Cell, value: CellValue, column: XlsxColumn) -> None:
    if value is None or value == "":
        return
    if isinstance(value, str):
        cell.value = _clean_text(value)
        cell.data_type = "s"
    else:
        cell.value = value
    if column.kind == "money":
        cell.number_format = MONEY_FORMAT
    elif column.kind == "datetime":
        cell.number_format = DATETIME_FORMAT
    elif column.kind == "integer":
        cell.number_format = INTEGER_FORMAT
    elif column.kind == "number":
        cell.number_format = NUMBER_FORMAT
    cell.alignment = WRAPPED_ALIGNMENT if column.kind == "long_text" else TOP_ALIGNMENT


def _column_width(column: XlsxColumn, rows: list[list[CellValue]], index: int) -> int:
    header_width = len(column.header) + 4
    if column.kind in FIXED_WIDTHS:
        return max(header_width, FIXED_WIDTHS[column.kind])
    longest_line = max(
        (
            len(line)
            for row in rows
            if row[index] is not None
            for line in str(row[index]).split("\n")
        ),
        default=0,
    )
    limit = LONG_TEXT_WIDTH if column.kind == "long_text" else MAX_WIDTH
    return max(MIN_WIDTH, header_width, min(longest_line + 2, limit))


def _write_sheet(sheet: Worksheet, content: XlsxSheet) -> None:
    sheet.title = content.title
    for index, column in enumerate(content.columns, start=1):
        cell = sheet.cell(row=1, column=index, value=column.header)
        cell.data_type = "s"
        cell.font = HEADER_FONT
        cell.alignment = TOP_ALIGNMENT
        sheet.column_dimensions[get_column_letter(index)].width = _column_width(
            column, content.rows, index - 1
        )
    for row_index, row in enumerate(content.rows, start=2):
        for column_index, (column, value) in enumerate(
            zip(content.columns, row, strict=True), start=1
        ):
            cell = sheet.cell(row=row_index, column=column_index)
            if isinstance(cell, Cell):
                _write_value(cell, value, column)
    last_column = get_column_letter(len(content.columns))
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:{last_column}{len(content.rows) + 1}"


class XlsxService:
    def render(self, sheets: list[XlsxSheet]) -> bytes:
        workbook = Workbook()
        default_sheet = workbook.active
        for index, content in enumerate(sheets):
            sheet = (
                default_sheet
                if index == 0 and isinstance(default_sheet, Worksheet)
                else workbook.create_sheet()
            )
            _write_sheet(sheet, content)
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
