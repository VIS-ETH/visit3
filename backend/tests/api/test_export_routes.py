import csv
import io
import re
import zipfile
from collections.abc import Awaitable, Callable
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from typing import Any
from urllib.parse import unquote
from uuid import UUID

import pytest
from httpx import AsyncClient, Response
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession
from sqlmodel import col, select

from app.models.company import KpCompanyProfile
from app.models.industry import Industry
from app.models.kp_event import (
    KpBookingCompanyDetails,
    KpEventBooking,
    KpEventBookingUpgradeWaitlist,
    NameTag,
)
from app.models.user import User
from app.repositories.kp_repository import KpRepository
from tests.api.conftest import (
    PNG_UPLOAD,
    company_profile_payload,
    first_member_id,
    kp_payload,
)

EVENT_NAME = 'KP "2026"/Süd'
SAFE_PREFIX = "KP -2026-Süd"
ASCII_PREFIX = "KP -2026-Sud"
OWN_COMPANY = "Acme AG"
OTHER_COMPANY = "Beta GmbH"
OTHER_COMPANY_EMAIL = "beta@example.com"
GENERAL_EMAIL = "info@acme.ch"
GENERAL_PHONE = "044 000 00 00"
INVOICE_ADDRESS = "Invoice street 1"
TEXT_ANSWER = "Two tables and a screen"
ALLOWED_UNTIL = date.today() + timedelta(days=3)
CSV_MEDIA_TYPE = "text/csv; charset=utf-8"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
COMPANY_WORKBOOK = "companies/download"
GERMAN_COMPANY_HEADERS = [
    "Unternehmen",
    "Marke",
    "Status",
    "Zone",
    "Stand-Nr.",
    "Standfläche (m²)",
    "Buchungs-Nr.",
    "Registriert am",
    "Bestätigt am",
    "Kontaktperson",
    "E-Mail Kontaktperson",
    "Telefon Kontaktperson",
    "Allgemeine E-Mail",
    "Allgemeine Telefonnummer",
    "Website",
    "Rechnungsempfänger",
    "Rechnungsadresse",
    "PLZ",
    "Ort",
    "Land",
    "MWST-Nr.",
    "Rechnungs-E-Mail",
    "Services",
    "Namensschilder",
    "Netto",
    "MWST",
    "Brutto",
    "Fehlende Angaben",
    "Beschreibung",
    "Branchen",
    "Sprachen",
    "Arbeitsorte",
    "Mitarbeitende Schweiz",
    "Mitarbeitende weltweit",
    "Praktika",
    "Teilzeitstellen",
    "Abschlussarbeiten",
    "Einstiegsstellen",
]
GERMAN_CONTACT_HEADERS = [
    "Unternehmen",
    "Status",
    "Zone",
    "Stand-Nr.",
    "Vorname",
    "Nachname",
    "E-Mail",
    "Telefon",
    "Kontaktperson",
]
ZIP_MEDIA_TYPE = "application/zip"
PDF_MEDIA_TYPE = "application/pdf"
DISPOSITION = re.compile(
    r'attachment; filename="(?P<ascii>[^"]+)"; filename\*=UTF-8\'\'(?P<encoded>\S+)'
)
UNSAFE_HEADER_CHARS = set('"\\/;\r\n')
TEXT_REQUIREMENT = {
    "type": "text",
    "name": "Booth layout",
    "description": "Describe the booth layout in a few sentences.",
}
FILE_REQUIREMENT = {
    "type": "file",
    "name": "Company logo",
    "description": "Upload the company logo for the booth wall.",
}
PRICE_HEADERS = ["net", "vat", "gross"]
BOOKING_HEADERS = [
    "booking_id",
    "company",
    "status",
    "zone",
    "booth_number",
    "booth_size_m2",
    "base_price",
    "services",
    "nametag_count",
    *PRICE_HEADERS,
    "missing_items",
]
EXPORT_HEADERS: dict[str, list[str]] = {
    "bookings/download": BOOKING_HEADERS,
    "waitlist-companies/download": [
        "company",
        "booking_id",
        "status",
        "current_zone",
        "current_booth_number",
        "target_zone",
        "priority_rank",
    ],
    "booked-services/download": [
        "company",
        "booking_id",
        "zone",
        "booth_number",
        "service",
        "quantity",
        "included_quantity",
        "charged_quantity",
        "unit_price",
        "total_charged_price",
        *PRICE_HEADERS,
    ],
    "nametags-data/download": [
        "name_tag_id",
        "booking_id",
        "company",
        "first_name",
        "last_name",
        "position",
    ],
    "service-requirements/download": [
        "company",
        "booking_id",
        "zone",
        "booth_number",
        "service",
        "requirement",
        "requirement_type",
        "uploaded",
        "text_value",
        "filename",
        "mime_type",
        "uploaded_at",
    ],
    "booth-zone-capacity/download": [
        "zone",
        "capacity",
        "booked_count",
        "remaining_capacity",
        "waitlist_demand",
        "booth_size_m2",
        "base_price",
        *PRICE_HEADERS,
    ],
    "registration-exceptions/download": ["company", "company_id", "allowed_until"],
}
EXPORT_FILENAMES: dict[str, str] = {
    "bookings/download": "-bookings-all.csv",
    "waitlist-companies/download": "-waitlist-companies.csv",
    "booked-services/download": "-booked-services.csv",
    "nametags-data/download": "-nametags-data.csv",
    "service-requirements/download": "-service-requirements-status.csv",
    "booth-zone-capacity/download": "-booth-zone-capacity.csv",
    "registration-exceptions/download": "-registration-exceptions.csv",
}


@dataclass(frozen=True)
class ExportWorld:
    event_id: str
    main_zone_id: str
    side_zone_id: str
    service_id: str
    booking_a: str
    booking_b: str
    name_tag_id: str
    headers: dict[str, str]


def read_csv(response: Response) -> tuple[list[str], list[dict[str, str]]]:
    reader = csv.DictReader(
        io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"
    )
    rows = list(reader)
    return list(reader.fieldnames or []), rows


def ok(response: Response) -> Response:
    assert response.status_code == 200, response.text
    return response


def disposition_names(response: Response) -> tuple[str, str]:
    match = DISPOSITION.fullmatch(response.headers["content-disposition"])
    assert match is not None
    return match["ascii"], unquote(match["encoded"])


async def export(client: AsyncClient, world: ExportWorld, suffix: str) -> Response:
    return await client.get(
        f"/api/kp/events/{world.event_id}/exports/{suffix}", headers=world.headers
    )


@pytest.fixture
async def other_company_user(create_user: Callable[..., Awaitable[User]]) -> User:
    return await create_user(
        email=OTHER_COMPANY_EMAIL, password=None, company_name=OTHER_COMPANY
    )


@pytest.fixture
async def export_world(
    client: AsyncClient,
    db_session: AsyncSession,
    company_user: User,
    company_headers: dict[str, str],
    staff_headers: dict[str, str],
    other_company_user: User,
    auth_headers: Callable[[User], Awaitable[dict[str, str]]],
    csrf_headers: dict[str, str],
) -> ExportWorld:
    other_headers = {**await auth_headers(other_company_user), **csrf_headers}
    event = ok(
        await client.post(
            "/api/kp/create", json=kp_payload(EVENT_NAME), headers=staff_headers
        )
    )
    event_id = event.json()["id"]
    zones = f"/api/kp/events/{event_id}/booth-zones"
    main_zone = ok(
        await client.post(
            zones,
            json={
                "name": "Main hall",
                "capacity": 5,
                "base_price": 10000,
                "booth_size": 12.5,
                "color": "#112233",
            },
            headers=staff_headers,
        )
    )
    side_zone = ok(
        await client.post(
            zones,
            json={
                "name": "Side hall",
                "capacity": 3,
                "base_price": 20000,
                "booth_size": 6.0,
                "color": "#445566",
            },
            headers=staff_headers,
        )
    )
    service = ok(
        await client.post(
            f"/api/kp/events/{event_id}/services",
            json={
                "name": "Electricity",
                "price": 5000,
                "max_quantity_per_booking": 2,
                "requirements": [TEXT_REQUIREMENT, FILE_REQUIREMENT],
            },
            headers=staff_headers,
        )
    )
    service_id = service.json()["id"]
    requirements = {
        requirement["name"]: requirement["id"]
        for requirement in service.json()["requirements"]
    }

    industry = Industry(name="Software")
    db_session.add(industry)
    await db_session.commit()
    booking_a = await register(
        client,
        company_headers,
        event_id,
        main_zone.json()["id"],
        service_id,
        profile=company_profile_payload(
            brand_name="Acme",
            description="<p>We build <strong>anvils</strong>.</p><p>Since 1900.</p>",
            general_email=GENERAL_EMAIL,
            general_phone=GENERAL_PHONE,
            employee_count_switzerland=42,
            employee_count_worldwide=99,
            offers_internships=True,
            languages=["ENGLISH", "GERMAN"],
            billing_street=INVOICE_ADDRESS,
            kp_contact_user_id=str(company_user.id),
            industry_ids=[str(industry.id)],
        ),
    )
    booking_b = await register(
        client, other_headers, event_id, side_zone.json()["id"], service_id
    )
    answers = f"/api/kp/booking-services/{booking_a['services'][0]['id']}/requirements"
    ok(
        await client.put(
            f"{answers}/{requirements['Booth layout']}/text",
            json={"text_value": TEXT_ANSWER},
            headers=company_headers,
        )
    )
    ok(
        await client.post(
            f"{answers}/{requirements['Company logo']}/file",
            files={"file": PNG_UPLOAD},
            headers=company_headers,
        )
    )
    kp_repository = KpRepository(db_session)
    await kp_repository.upsert_registration_exception(
        event_id=UUID(event_id),
        company_id=UUID(str(other_company_user.company_id)),
        allowed_until=ALLOWED_UNTIL,
    )
    name_tags = [
        NameTag(
            booking_id=UUID(booking_a["id"]),
            first_name="Ada",
            last_name="Lovelace",
            position="Engineer",
        ),
        NameTag(
            booking_id=UUID(booking_a["id"]),
            first_name="Grace",
            last_name="Hopper",
            position="Admiral",
        ),
        NameTag(
            booking_id=UUID(booking_b["id"]),
            first_name="Alan",
            last_name="Turing",
            position="Analyst",
        ),
    ]
    db_session.add_all(name_tags)
    await db_session.commit()
    name_tag_id = str(name_tags[0].id)
    db_session.expire_all()

    ok(
        await client.patch(
            f"/api/kp/bookings/{booking_b['id']}/status",
            json={"status": "CANCELLED"},
            headers=other_headers,
        )
    )

    db_session.add_all(
        [
            KpEventBookingUpgradeWaitlist(
                booking_id=UUID(booking_a["id"]),
                target_booth_zone_id=UUID(side_zone.json()["id"]),
                priority_rank=1,
            ),
            KpEventBookingUpgradeWaitlist(
                booking_id=UUID(booking_b["id"]),
                target_booth_zone_id=UUID(main_zone.json()["id"]),
                priority_rank=1,
            ),
        ]
    )
    await db_session.commit()
    ok(
        await client.post(
            f"/api/kp/events/{event_id}/exports/nametags/background",
            files={"file": PNG_UPLOAD},
            headers=staff_headers,
        )
    )

    return ExportWorld(
        event_id=event_id,
        main_zone_id=main_zone.json()["id"],
        side_zone_id=side_zone.json()["id"],
        service_id=service_id,
        booking_a=booking_a["id"],
        booking_b=booking_b["id"],
        name_tag_id=name_tag_id,
        headers=staff_headers,
    )


async def register(
    client: AsyncClient,
    headers: dict[str, str],
    event_id: str,
    booth_zone_id: str,
    service_id: str,
    profile: dict[str, Any] | None = None,
) -> dict[str, Any]:
    ok(
        await client.put(
            "/api/company/me/profile",
            json=profile
            if profile is not None
            else company_profile_payload(
                kp_contact_user_id=await first_member_id(client, headers)
            ),
            headers=headers,
        )
    )
    response = ok(
        await client.post(
            f"/api/kp/events/{event_id}/bookings/register",
            json={
                "booth_zone_id": booth_zone_id,
                "services": [{"service_id": service_id, "quantity": 1}],
                "confirm_profile": True,
            },
            headers=headers,
        )
    )
    return response.json()


@pytest.mark.parametrize("suffix", sorted(EXPORT_FILENAMES))
async def test_csv_export_is_a_hardened_download(
    client: AsyncClient, export_world: ExportWorld, suffix: str
):
    response = await export(client, export_world, suffix)

    ascii_name, encoded_name = disposition_names(response)
    assert response.status_code == 200
    assert response.headers["content-type"] == CSV_MEDIA_TYPE
    assert ascii_name == f"{ASCII_PREFIX}{EXPORT_FILENAMES[suffix]}"
    assert encoded_name == f"{SAFE_PREFIX}{EXPORT_FILENAMES[suffix]}"
    assert not UNSAFE_HEADER_CHARS & set(ascii_name)


@pytest.mark.parametrize("suffix", sorted(EXPORT_HEADERS))
async def test_csv_export_header_row(
    client: AsyncClient, export_world: ExportWorld, suffix: str
):
    fieldnames, _ = read_csv(await export(client, export_world, suffix))

    assert fieldnames == EXPORT_HEADERS[suffix]


async def test_bookings_export_keeps_cancelled_bookings(
    client: AsyncClient, export_world: ExportWorld
):
    _, rows = read_csv(await export(client, export_world, "bookings/download"))

    by_company = {row["company"]: row for row in rows}
    assert by_company[OWN_COMPANY]["status"] == "REGISTERED"
    assert by_company[OWN_COMPANY]["zone"] == "Main hall"
    assert by_company[OWN_COMPANY]["base_price"] == "100.00"
    assert by_company[OWN_COMPANY]["services"] == "Electricity x1"
    assert by_company[OWN_COMPANY]["nametag_count"] == "2"
    assert by_company[OWN_COMPANY]["net"] == "150.00"
    assert by_company[OWN_COMPANY]["vat"] == "12.15"
    assert by_company[OWN_COMPANY]["gross"] == "162.15"
    assert by_company[OTHER_COMPANY]["status"] == "CANCELLED"


async def strip_company_profile(db_session: AsyncSession, booking_id: str) -> None:
    snapshot = (
        await db_session.execute(
            select(KpBookingCompanyDetails).where(
                col(KpBookingCompanyDetails.booking_id) == UUID(booking_id)
            )
        )
    ).scalar_one()
    booking = (
        await db_session.execute(
            select(KpEventBooking).where(col(KpEventBooking.id) == UUID(booking_id))
        )
    ).scalar_one()
    profile = (
        await db_session.execute(
            select(KpCompanyProfile).where(
                col(KpCompanyProfile.company_id) == booking.company_id
            )
        )
    ).scalar_one()
    profile.billing_city = ""
    await db_session.delete(snapshot)
    db_session.add(profile)
    await db_session.commit()


async def test_bookings_export_lists_missing_items(
    client: AsyncClient, db_session: AsyncSession, export_world: ExportWorld
):
    await strip_company_profile(db_session, export_world.booking_b)

    _, rows = read_csv(await export(client, export_world, "bookings/download"))

    by_company = {row["company"]: row for row in rows}
    assert by_company[OWN_COMPANY]["missing_items"] == ""
    missing = by_company[OTHER_COMPANY]["missing_items"].split("; ")
    assert "company_profile" in missing
    assert "billing_address" in missing
    assert sum(item.startswith("requirement:") for item in missing) == 2


def open_workbook(response: Response) -> Workbook:
    return load_workbook(io.BytesIO(response.content))


def sheet_rows(workbook: Workbook, title: str) -> list[dict[str, Any]]:
    rows = [list(row) for row in workbook[title].iter_rows(values_only=True)]
    return [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]


async def test_company_workbook_is_an_xlsx_download(
    client: AsyncClient, export_world: ExportWorld
):
    response = await export(client, export_world, COMPANY_WORKBOOK)

    ascii_name, encoded_name = disposition_names(response)
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert encoded_name == f"{SAFE_PREFIX} Unternehmen {date.today().isoformat()}.xlsx"
    assert ascii_name == f"{ASCII_PREFIX} Unternehmen {date.today().isoformat()}.xlsx"
    assert open_workbook(response).sheetnames == ["Unternehmen", "Kontakte"]


async def test_company_workbook_headers_are_human(
    client: AsyncClient, export_world: ExportWorld
):
    workbook = open_workbook(await export(client, export_world, COMPANY_WORKBOOK))

    firms = workbook["Unternehmen"]
    assert [cell.value for cell in firms[1]] == GERMAN_COMPANY_HEADERS
    assert [cell.value for cell in workbook["Kontakte"][1]] == GERMAN_CONTACT_HEADERS
    assert firms.freeze_panes == "B2"
    assert firms.auto_filter.ref == f"A1:AL{firms.max_row}"


async def test_company_workbook_follows_the_staff_language(
    client: AsyncClient, export_world: ExportWorld
):
    response = await export(client, export_world, f"{COMPANY_WORKBOOK}?language=en")

    _, encoded_name = disposition_names(response)
    workbook = open_workbook(response)
    rows = sheet_rows(workbook, "Companies")
    assert encoded_name == f"{SAFE_PREFIX} Companies {date.today().isoformat()}.xlsx"
    assert workbook.sheetnames == ["Companies", "Contacts"]
    assert [cell.value for cell in workbook["Companies"][1]][:4] == [
        "Company",
        "Brand",
        "Status",
        "Zone",
    ]
    own = next(row for row in rows if row["Company"] == OWN_COMPANY)
    assert own["Status"] == "Registered"
    assert own["Internships"] == "Yes"
    assert own["Languages"] == "English, German"


async def test_company_workbook_rows_carry_typed_values(
    client: AsyncClient, export_world: ExportWorld
):
    workbook = open_workbook(await export(client, export_world, COMPANY_WORKBOOK))
    rows = sheet_rows(workbook, "Unternehmen")

    assert [row["Unternehmen"] for row in rows] == [OWN_COMPANY, OTHER_COMPANY]
    own = rows[0]
    assert own["Marke"] == "Acme"
    assert own["Status"] == "Registriert"
    assert own["Zone"] == "Main hall"
    assert own["Standfläche (m²)"] == 12.5
    assert isinstance(own["Buchungs-Nr."], int)
    assert isinstance(own["Registriert am"], datetime)
    assert own["Bestätigt am"] is None
    assert own["Kontaktperson"] == "Test User"
    assert own["E-Mail Kontaktperson"] == "company@example.com"
    assert own["Allgemeine E-Mail"] == GENERAL_EMAIL
    assert own["Allgemeine Telefonnummer"] == GENERAL_PHONE
    assert own["Rechnungsadresse"] == f"{INVOICE_ADDRESS} 1"
    assert own["PLZ"] == "8000"
    assert own["Land"] == "CH"
    assert own["Services"] == "Electricity × 1"
    assert own["Namensschilder"] == 2
    assert own["Netto"] == 150.0
    assert own["Brutto"] == 162.15
    assert own["Beschreibung"] == "We build anvils.\nSince 1900."
    assert own["Branchen"] == "Software"
    assert own["Sprachen"] == "Englisch, Deutsch"
    assert own["Mitarbeitende Schweiz"] == 42
    assert own["Praktika"] == "Ja"
    assert own["Abschlussarbeiten"] == "Nein"
    assert rows[1]["Status"] == "Storniert"
    assert workbook["Unternehmen"]["Y2"].number_format == '"CHF" #,##0.00'
    assert workbook["Unternehmen"]["H2"].number_format == "DD.MM.YYYY HH:MM"


async def test_company_workbook_names_what_is_missing(
    client: AsyncClient, db_session: AsyncSession, export_world: ExportWorld
):
    await strip_company_profile(db_session, export_world.booking_b)

    rows = sheet_rows(
        open_workbook(await export(client, export_world, COMPANY_WORKBOOK)),
        "Unternehmen",
    )

    other = next(row for row in rows if row["Unternehmen"] == OTHER_COMPANY)
    missing = other["Fehlende Angaben"].split("\n")
    assert "Unternehmensprofil" in missing
    assert "Rechnungsadresse" in missing
    assert "Electricity · Booth layout" in missing
    assert other["Marke"] is None
    assert other["Beschreibung"] is None
    assert other["Praktika"] is None


async def test_company_workbook_keeps_company_text_inert(
    client: AsyncClient, db_session: AsyncSession, export_world: ExportWorld
):
    hostile = '=HYPERLINK("https://evil.test","Klick")'
    snapshot = (
        await db_session.execute(
            select(KpBookingCompanyDetails).where(
                col(KpBookingCompanyDetails.booking_id) == UUID(export_world.booking_a)
            )
        )
    ).scalar_one()
    snapshot.brand_name = hostile
    snapshot.description = "<p>Grüezi Zürich 🤖</p><p>@SUM(A1)</p>"
    db_session.add(snapshot)
    await db_session.commit()

    sheet = open_workbook(await export(client, export_world, COMPANY_WORKBOOK))[
        "Unternehmen"
    ]

    assert sheet["B2"].value == hostile
    assert sheet["B2"].data_type == "s"
    assert sheet["AC2"].value == "Grüezi Zürich 🤖\n@SUM(A1)"
    assert sheet["AC2"].data_type == "s"


async def test_company_workbook_lists_the_people_of_active_bookings(
    client: AsyncClient, export_world: ExportWorld
):
    rows = sheet_rows(
        open_workbook(await export(client, export_world, COMPANY_WORKBOOK)),
        "Kontakte",
    )

    assert rows == [
        {
            "Unternehmen": OWN_COMPANY,
            "Status": "Registriert",
            "Zone": "Main hall",
            "Stand-Nr.": None,
            "Vorname": "Test",
            "Nachname": "User",
            "E-Mail": "company@example.com",
            "Telefon": None,
            "Kontaktperson": "Ja",
        }
    ]


async def test_booth_zone_capacity_export_excludes_cancelled_bookings(
    client: AsyncClient, export_world: ExportWorld
):
    _, rows = read_csv(
        await export(client, export_world, "booth-zone-capacity/download")
    )

    by_zone = {row["zone"]: row for row in rows}
    assert by_zone["Main hall"]["booked_count"] == "1"
    assert by_zone["Main hall"]["remaining_capacity"] == "4"
    assert by_zone["Main hall"]["waitlist_demand"] == "0"
    assert by_zone["Side hall"]["booked_count"] == "0"
    assert by_zone["Side hall"]["remaining_capacity"] == "3"
    assert by_zone["Side hall"]["waitlist_demand"] == "1"
    assert by_zone["Main hall"]["net"] == "100.00"
    assert by_zone["Main hall"]["vat"] == "8.10"
    assert by_zone["Main hall"]["gross"] == "108.10"
    assert by_zone["Side hall"]["vat"] == "16.20"


async def test_waitlist_export_excludes_cancelled_bookings(
    client: AsyncClient, export_world: ExportWorld
):
    _, rows = read_csv(
        await export(client, export_world, "waitlist-companies/download")
    )

    assert [row["company"] for row in rows] == [OWN_COMPANY]
    assert rows[0]["current_zone"] == "Main hall"
    assert rows[0]["target_zone"] == "Side hall"
    assert rows[0]["priority_rank"] == "1"


async def test_booked_services_export_rows(
    client: AsyncClient, export_world: ExportWorld
):
    _, rows = read_csv(await export(client, export_world, "booked-services/download"))

    by_company = {row["company"]: row for row in rows}
    assert sorted(by_company) == [OWN_COMPANY, OTHER_COMPANY]
    assert by_company[OWN_COMPANY]["service"] == "Electricity"
    assert by_company[OWN_COMPANY]["charged_quantity"] == "1"
    assert by_company[OWN_COMPANY]["unit_price"] == "50.00"
    assert by_company[OWN_COMPANY]["total_charged_price"] == "50.00"
    assert by_company[OWN_COMPANY]["net"] == "50.00"
    assert by_company[OWN_COMPANY]["vat"] == "4.05"
    assert by_company[OWN_COMPANY]["gross"] == "54.05"


async def test_service_requirements_export_rows(
    client: AsyncClient, export_world: ExportWorld
):
    _, rows = read_csv(
        await export(client, export_world, "service-requirements/download")
    )

    own = {row["requirement"]: row for row in rows if row["company"] == OWN_COMPANY}
    other = {row["requirement"]: row for row in rows if row["company"] == OTHER_COMPANY}
    assert own["Booth layout"]["uploaded"] == "yes"
    assert own["Booth layout"]["text_value"] == TEXT_ANSWER
    assert own["Company logo"]["uploaded"] == "yes"
    assert own["Company logo"]["filename"] == PNG_UPLOAD[0]
    assert other["Booth layout"]["uploaded"] == "no"


async def test_nametags_data_export_rows(
    client: AsyncClient, export_world: ExportWorld
):
    _, rows = read_csv(await export(client, export_world, "nametags-data/download"))

    assert [(row["last_name"], row["company"]) for row in rows] == [
        ("Hopper", OWN_COMPANY),
        ("Lovelace", OWN_COMPANY),
        ("Turing", OTHER_COMPANY),
    ]


async def test_registration_exceptions_export_rows(
    client: AsyncClient, export_world: ExportWorld
):
    _, rows = read_csv(
        await export(client, export_world, "registration-exceptions/download")
    )

    assert [row["company"] for row in rows] == [OTHER_COMPANY]
    assert rows[0]["allowed_until"] == str(ALLOWED_UNTIL)


async def test_bookings_by_zone_zip_has_one_entry_per_zone(
    client: AsyncClient, export_world: ExportWorld
):
    response = await export(client, export_world, "bookings/by-zone/download")

    ascii_name, encoded_name = disposition_names(response)
    archive = zipfile.ZipFile(io.BytesIO(response.content))
    main = csv.DictReader(
        io.StringIO(archive.read("main-hall.csv").decode("utf-8-sig")), delimiter=";"
    )
    side = csv.DictReader(
        io.StringIO(archive.read("side-hall.csv").decode("utf-8-sig")), delimiter=";"
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == ZIP_MEDIA_TYPE
    assert ascii_name == f"{ASCII_PREFIX}-bookings-by-zone.zip"
    assert encoded_name == f"{SAFE_PREFIX}-bookings-by-zone.zip"
    assert archive.namelist() == ["main-hall.csv", "side-hall.csv"]
    assert main.fieldnames == BOOKING_HEADERS
    assert [(row["company"], row["net"], row["vat"]) for row in main] == [
        (OWN_COMPANY, "150.00", "12.15")
    ]
    assert [row["company"] for row in side] == [OTHER_COMPANY]


async def test_nametag_export_targets(client: AsyncClient, export_world: ExportWorld):
    response = await client.get(
        f"/api/kp/events/{export_world.event_id}/exports/nametags/targets",
        headers=export_world.headers,
    )

    body = response.json()
    assert response.status_code == 200
    assert [company["company_name"] for company in body["companies"]] == [
        OWN_COMPANY,
        OTHER_COMPANY,
    ]
    assert [person["last_name"] for person in body["people"]] == [
        "Hopper",
        "Lovelace",
        "Turing",
    ]


async def test_nametag_background_is_stored_and_readable(
    client: AsyncClient, export_world: ExportWorld
):
    response = await client.get(
        f"/api/kp/events/{export_world.event_id}/exports/nametags/background",
        headers=export_world.headers,
    )

    assert response.status_code == 200
    assert response.json()["stored_file"]["original_filename"] == PNG_UPLOAD[0]


async def test_event_nametags_pdf_download(
    client: AsyncClient, export_world: ExportWorld
):
    response = await export(client, export_world, "nametags/download")

    ascii_name, _ = disposition_names(response)
    assert response.status_code == 200
    assert response.headers["content-type"] == PDF_MEDIA_TYPE
    assert response.content.startswith(b"%PDF")
    assert ascii_name == f"{ASCII_PREFIX}-nametags.pdf"


async def test_booking_nametags_pdf_download(
    client: AsyncClient, export_world: ExportWorld
):
    response = await client.get(
        f"/api/kp/bookings/{export_world.booking_a}/nametags/download",
        headers=export_world.headers,
    )

    ascii_name, _ = disposition_names(response)
    assert response.status_code == 200
    assert response.headers["content-type"] == PDF_MEDIA_TYPE
    assert response.content.startswith(b"%PDF")
    assert ascii_name == f"{ASCII_PREFIX}-Acme AG-nametags.pdf"


async def test_single_nametag_pdf_download(
    client: AsyncClient, export_world: ExportWorld
):
    response = await client.get(
        f"/api/kp/nametags/{export_world.name_tag_id}/download",
        headers=export_world.headers,
    )

    ascii_name, _ = disposition_names(response)
    assert response.status_code == 200
    assert response.headers["content-type"] == PDF_MEDIA_TYPE
    assert response.content.startswith(b"%PDF")
    assert ascii_name == f"{ASCII_PREFIX}-Ada-Lovelace-nametag.pdf"
