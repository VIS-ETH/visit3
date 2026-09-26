import csv
import io
import zipfile
from collections.abc import Sequence
from datetime import date, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.core.exceptions import ExportRenderTimeout
from app.models.company import Company
from app.models.kp_event import (
    KpBookingCompanyDetails,
    KpBookingStatus,
    KpCompanyLanguage,
    KpEvent,
    KpEventBooking,
    KpEventBookingUpgradeWaitlist,
    KpEventBoothZone,
    NameTag,
)
from app.models.user import User
from app.services.csv_service import CsvService
from app.services.export_service import (
    NAMETAG_TEMPLATE_NAME,
    ExportLanguage,
    ExportService,
)
from app.services.typst_runner import TypstRenderAborted
from app.services.xlsx_service import XlsxService


def make_event(name: str = "Kontaktparty", vat_rate_permille: int = 81) -> KpEvent:
    today = date.today()
    return KpEvent(
        id=uuid4(),
        name=name,
        vat_rate_permille=vat_rate_permille,
        registration_open=today - timedelta(days=10),
        registration_end=today - timedelta(days=5),
        finalization_deadline=today - timedelta(days=4),
        nametags_deadline=today - timedelta(days=3),
        event_date=today + timedelta(days=10),
    )


def make_zone(
    *,
    event_id,
    name: str,
    capacity: int = 5,
    color: str = "#000000",
    base_price: int = 0,
) -> KpEventBoothZone:
    return KpEventBoothZone(
        id=uuid4(),
        event_id=event_id,
        name=name,
        description="",
        color=color,
        capacity=capacity,
        base_price=base_price,
    )


def make_name_tag(*, booking_id, first_name: str = "Ada") -> NameTag:
    return NameTag(
        id=uuid4(),
        booking_id=booking_id,
        first_name=first_name,
        last_name="Lovelace",
        position="Engineer",
    )


def make_waitlist_entry(*, booking_id, target_zone) -> KpEventBookingUpgradeWaitlist:
    entry = KpEventBookingUpgradeWaitlist(
        id=uuid4(),
        booking_id=booking_id,
        target_booth_zone_id=target_zone.id,
        priority_rank=1,
    )
    entry.target_booth_zone = target_zone
    return entry


def make_export_booking(
    *,
    event: KpEvent,
    zone: KpEventBoothZone,
    company_name: str = "Acme AG",
    status: KpBookingStatus = KpBookingStatus.REGISTERED,
    booth_nr: int | None = None,
    nametag_count: int = 0,
    waitlist_targets: Sequence[KpEventBoothZone] = (),
) -> KpEventBooking:
    booking = KpEventBooking(
        id=uuid4(),
        event_id=event.id,
        company_id=uuid4(),
        booth_zone_id=zone.id,
        status=status,
        booth_nr=booth_nr,
    )
    booking.company = Company(id=booking.company_id, name=company_name)
    booking.booth_zone = zone
    booking.name_tags = [
        make_name_tag(booking_id=booking.id) for _ in range(nametag_count)
    ]
    booking.upgrade_waitlist_entries = [
        make_waitlist_entry(booking_id=booking.id, target_zone=target)
        for target in waitlist_targets
    ]
    booking.services = []
    booking.company_details = None
    return booking


def make_export_service(kp_repo, storage_service, staff_user) -> ExportService:
    return ExportService(
        kp_repository=kp_repo,
        storage_service=storage_service,
        pdf_service=Mock(),
        csv_service=CsvService(),
        xlsx_service=XlsxService(),
        current_user=staff_user,
    )


def read_csv_rows(content: bytes) -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig")), delimiter=";"))


class FakePdfService:
    def __init__(self) -> None:
        self.rendered_data: dict[str, object] | None = None
        self.rendered_filename: str | None = None

    async def render(
        self,
        template_name: str,
        data: dict[str, object],
        filename: str,
        root: str | None = None,
        template_dir: Path | None = None,
    ):
        assert root is not None
        root_path = Path(root)
        assert template_dir == root_path
        assert template_name == NAMETAG_TEMPLATE_NAME
        assert data["background_path"] == "background.png"
        assert (root_path / "background.png").read_bytes() == b"\x89PNG\r\n"
        self.rendered_data = data
        self.rendered_filename = filename
        return b"%PDF", filename


async def test_nametag_pdf_render_uses_restricted_workspace_and_json_data():
    pdf_service = FakePdfService()
    service = ExportService(
        kp_repository=Mock(),
        storage_service=Mock(),
        pdf_service=pdf_service,
        csv_service=CsvService(),
        xlsx_service=XlsxService(),
        current_user=Mock(),
    )
    name_tag = SimpleNamespace(
        first_name='#panic("boom")',
        last_name="Person",
        position="#image('/etc/passwd')",
        booking=SimpleNamespace(company=SimpleNamespace(name="=Company")),
    )

    export = await service._render_nametags_pdf(
        b"\x89PNG\r\n",
        "image/png",
        [name_tag],
        '../"nametag\r.pdf',
        1,
    )

    assert export.content == b"%PDF"
    assert export.filename == "nametag.pdf"
    assert pdf_service.rendered_data is not None
    assert pdf_service.rendered_data["tags"] == [
        {
            "full_name": '#panic("boom") Person',
            "position": "#image('/etc/passwd')",
            "company": "=Company",
        }
    ]


class OverloadedPdfService:
    async def render(self, *args: object, **kwargs: object):
        raise TypstRenderAborted("timeout")


async def test_a_nametag_export_that_renders_too_long_is_reported():
    service = ExportService(
        kp_repository=Mock(),
        storage_service=Mock(),
        pdf_service=OverloadedPdfService(),
        csv_service=CsvService(),
        xlsx_service=XlsxService(),
        current_user=Mock(),
    )
    name_tag = SimpleNamespace(
        first_name="Ada",
        last_name="Lovelace",
        position="Engineer",
        booking=SimpleNamespace(company=SimpleNamespace(name="Acme")),
    )

    with pytest.raises(ExportRenderTimeout) as error:
        await service._render_nametags_pdf(
            b"\x89PNG\r\n", "image/png", [name_tag], "nametags.pdf", 2
        )

    assert error.value.code == "error.export_render_timeout"
    assert error.value.status_code == 503


async def test_list_nametag_export_targets_sorts_unassigned_booths_last(
    kp_repo,
    storage_service,
    staff_user,
):
    event = make_event()
    zone = make_zone(event_id=event.id, name="Main hall")
    unassigned = make_export_booking(event=event, zone=zone, nametag_count=1)
    assigned = make_export_booking(event=event, zone=zone, booth_nr=3, nametag_count=1)
    kp_repo.get_by_id.return_value = event
    kp_repo.list_bookings_for_event.return_value = [unassigned, assigned]
    service = make_export_service(kp_repo, storage_service, staff_user)

    result = await service.list_nametag_export_targets(event.id)

    assert [company.booth_nr for company in result.companies] == [3, None]


async def test_booth_zone_capacity_export_ignores_cancelled_bookings(
    kp_repo,
    storage_service,
    staff_user,
):
    event = make_event()
    zone = make_zone(event_id=event.id, name="Main hall", capacity=5)
    target_zone = make_zone(event_id=event.id, name="Gold", color="#111111")
    active = make_export_booking(event=event, zone=zone, waitlist_targets=[target_zone])
    cancelled = make_export_booking(
        event=event,
        zone=zone,
        company_name="Gone AG",
        status=KpBookingStatus.CANCELLED,
        waitlist_targets=[target_zone],
    )
    kp_repo.get_by_id.return_value = event
    kp_repo.list_bookings_for_event.return_value = [active, cancelled]
    kp_repo.list_booth_zones.return_value = [zone, target_zone]
    service = make_export_service(kp_repo, storage_service, staff_user)

    export = await service.export_booth_zone_capacity_csv(event.id)

    rows = read_csv_rows(export.content)
    assert [
        (row["zone"], row["booked_count"], row["remaining_capacity"]) for row in rows
    ] == [("Main hall", "1", "4"), ("Gold", "0", "5")]
    assert [row["waitlist_demand"] for row in rows] == ["0", "1"]


async def test_waitlist_export_skips_cancelled_bookings(
    kp_repo,
    storage_service,
    staff_user,
):
    event = make_event()
    zone = make_zone(event_id=event.id, name="Main hall")
    target_zone = make_zone(event_id=event.id, name="Gold", color="#111111")
    active = make_export_booking(event=event, zone=zone, waitlist_targets=[target_zone])
    cancelled = make_export_booking(
        event=event,
        zone=zone,
        company_name="Gone AG",
        status=KpBookingStatus.CANCELLED,
        waitlist_targets=[target_zone],
    )
    kp_repo.get_by_id.return_value = event
    kp_repo.list_bookings_for_event.return_value = [active, cancelled]
    service = make_export_service(kp_repo, storage_service, staff_user)

    export = await service.export_waitlist_companies_csv(event.id)

    assert [row["company"] for row in read_csv_rows(export.content)] == ["Acme AG"]


async def test_bookings_by_zone_zip_deduplicates_colliding_entry_names(
    kp_repo,
    storage_service,
    staff_user,
):
    event = make_event()
    first_zone = make_zone(event_id=event.id, name="Hall A/B")
    second_zone = make_zone(event_id=event.id, name="Hall A-B", color="#111111")
    third_zone = make_zone(event_id=event.id, name="Hall A B", color="#222222")
    kp_repo.get_by_id.return_value = event
    kp_repo.list_bookings_for_event.return_value = [
        make_export_booking(event=event, zone=first_zone)
    ]
    kp_repo.list_booth_zones.return_value = [first_zone, second_zone, third_zone]
    service = make_export_service(kp_repo, storage_service, staff_user)

    export = await service.export_bookings_by_zone_zip(event.id)

    with zipfile.ZipFile(io.BytesIO(export.content)) as archive:
        assert archive.namelist() == [
            "hall-a-b.csv",
            "hall-a-b-2.csv",
            "hall-a-b-3.csv",
        ]
        assert [
            row["company"] for row in read_csv_rows(archive.read("hall-a-b.csv"))
        ] == ["Acme AG"]
        assert read_csv_rows(archive.read("hall-a-b-2.csv")) == []


async def test_company_workbook_lists_only_people_of_active_bookings(
    kp_repo,
    storage_service,
    staff_user,
):
    event = make_event()
    zone = make_zone(event_id=event.id, name="Main hall")
    active = make_export_booking(event=event, zone=zone)
    cancelled = make_export_booking(
        event=event,
        zone=zone,
        company_name="Gone AG",
        status=KpBookingStatus.CANCELLED,
    )
    for booking, email in ((active, "ada@acme.test"), (cancelled, "gone@gone.test")):
        booking.company.users = [
            User(email=email, first_name="Ada", company_id=booking.company_id)
        ]
    kp_repo.get_by_id.return_value = event
    kp_repo.list_bookings_for_event.return_value = [active, cancelled]
    service = make_export_service(kp_repo, storage_service, staff_user)

    export = await service.export_company_workbook(event.id, ExportLanguage.DE)

    contacts = load_workbook(io.BytesIO(export.content))["Kontakte"]
    assert [row[6] for row in contacts.iter_rows(min_row=2, values_only=True)] == [
        "ada@acme.test"
    ]


async def test_company_workbook_names_languages_in_the_staff_language(
    kp_repo,
    storage_service,
    staff_user,
):
    event = make_event()
    zone = make_zone(event_id=event.id, name="Main hall")
    booking = make_export_booking(event=event, zone=zone)
    booking.company.users = []
    booking.company_details = KpBookingCompanyDetails(
        id=uuid4(),
        booking_id=booking.id,
        languages=[KpCompanyLanguage.ENGLISH, KpCompanyLanguage.GERMAN],
    )
    kp_repo.get_by_id.return_value = event
    kp_repo.list_bookings_for_event.return_value = [booking]
    service = make_export_service(kp_repo, storage_service, staff_user)

    german = await service.export_company_workbook(event.id, ExportLanguage.DE)
    english = await service.export_company_workbook(event.id, ExportLanguage.EN)

    assert load_workbook(io.BytesIO(german.content))["Unternehmen"]["AE2"].value == (
        "Englisch, Deutsch"
    )
    assert load_workbook(io.BytesIO(english.content))["Companies"]["AE2"].value == (
        "English, German"
    )


async def test_booking_export_prices_use_the_event_vat_rate(
    kp_repo,
    storage_service,
    staff_user,
):
    event = make_event(vat_rate_permille=77)
    zone = make_zone(event_id=event.id, name="Main hall", base_price=1999)
    booking = make_export_booking(event=event, zone=zone)
    kp_repo.get_by_id.return_value = event
    kp_repo.list_bookings_for_event.return_value = [booking]
    service = make_export_service(kp_repo, storage_service, staff_user)

    export = await service.export_bookings_csv(event.id)

    rows = read_csv_rows(export.content)
    assert [(row["net"], row["vat"], row["gross"]) for row in rows] == [
        ("19.99", "1.54", "21.53")
    ]


async def test_capacity_export_prices_use_the_event_vat_rate(
    kp_repo,
    storage_service,
    staff_user,
):
    event = make_event(vat_rate_permille=26)
    zone = make_zone(event_id=event.id, name="Main hall", base_price=12345)
    kp_repo.get_by_id.return_value = event
    kp_repo.list_bookings_for_event.return_value = []
    kp_repo.list_booth_zones.return_value = [zone]
    service = make_export_service(kp_repo, storage_service, staff_user)

    export = await service.export_booth_zone_capacity_csv(event.id)

    rows = read_csv_rows(export.content)
    assert [
        (row["base_price"], row["net"], row["vat"], row["gross"]) for row in rows
    ] == [("123.45", "123.45", "3.21", "126.66")]
