import io
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.services.xlsx_service import XlsxColumn, XlsxService, XlsxSheet

COLUMNS = [
    XlsxColumn("Firma", "text"),
    XlsxColumn("Stand-Nr.", "integer"),
    XlsxColumn("Brutto", "money"),
    XlsxColumn("Registriert am", "datetime"),
    XlsxColumn("Praktika", "text"),
    XlsxColumn("Beschreibung", "long_text"),
]


def render(rows: list[list[object]], title: str = "Firmen") -> Worksheet:
    content = XlsxService().render([XlsxSheet(title, COLUMNS, rows)])
    return load_workbook(io.BytesIO(content))[title]


def test_the_header_row_is_bold_frozen_and_filterable():
    sheet = render(
        [["Acme AG", 12, 1234.5, datetime(2026, 9, 1, 14, 30), "Ja", "Text"]]
    )

    assert [cell.value for cell in sheet[1]] == [column.header for column in COLUMNS]
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet.freeze_panes == "B2"
    assert sheet.auto_filter.ref == "A1:F2"


def test_values_keep_their_types_and_formats():
    sheet = render(
        [["Acme AG", 12, 1234.5, datetime(2026, 9, 1, 14, 30), "Ja", "Text"]]
    )
    company, booth, gross, registered, internships, description = sheet[2]

    assert company.value == "Acme AG"
    assert booth.value == 12
    assert booth.data_type == "n"
    assert gross.value == 1234.5
    assert "CHF" in gross.number_format
    assert gross.number_format.endswith("0.00")
    assert registered.value == datetime(2026, 9, 1, 14, 30)
    assert registered.number_format == "DD.MM.YYYY HH:MM"
    assert internships.value == "Ja"
    assert description.alignment.wrap_text is True


def test_formula_like_text_stays_an_inert_string():
    hostile = [
        '=HYPERLINK("https://evil.test","click")',
        "+41 44 123 45 67",
        "-1+2",
        "@SUM(A1:A2)",
        "=1+1",
    ]
    sheet = render([[text, None, None, None, None, None] for text in hostile])

    cells = [sheet.cell(row=row, column=1) for row in range(2, 2 + len(hostile))]
    assert [cell.value for cell in cells] == hostile
    assert {cell.data_type for cell in cells} == {"s"}


def test_umlauts_emoji_and_line_breaks_survive():
    text = "Zürich Grüezi 🤖\nZweite Zeile"
    sheet = render([["Müller & Söhne AG", None, None, None, None, text]])

    assert sheet["A2"].value == "Müller & Söhne AG"
    assert sheet["F2"].value == text


def test_control_characters_are_removed():
    sheet = render([["Acme\x00\x0b AG", None, None, None, None, "a\x07b"]])

    assert sheet["A2"].value == "Acme AG"
    assert sheet["F2"].value == "ab"


def test_empty_values_leave_empty_cells():
    sheet = render([["Acme AG", None, None, None, "", None]])

    assert [cell.value for cell in sheet[2]] == [
        "Acme AG",
        None,
        None,
        None,
        None,
        None,
    ]


def test_columns_are_wide_enough_but_capped():
    sheet = render([["A" * 200, 1, 1.0, None, "Ja", "B" * 500]])

    assert 10 <= sheet.column_dimensions["A"].width <= 60
    assert sheet.column_dimensions["B"].width >= len("Stand-Nr.")
    assert sheet.column_dimensions["F"].width <= 80


def test_every_sheet_is_written_in_order():
    content = XlsxService().render(
        [
            XlsxSheet("Firmen", COLUMNS[:1], [["Acme AG"]]),
            XlsxSheet("Kontakte", COLUMNS[:1], []),
        ]
    )

    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["Firmen", "Kontakte"]
    assert workbook["Kontakte"].auto_filter.ref == "A1:A1"
